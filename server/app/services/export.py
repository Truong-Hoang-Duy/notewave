import io
from zoneinfo import ZoneInfo

from docx import Document
from docx.shared import Pt, RGBColor

from app.models.session import SessionRead
from app.services.transcript import format_timestamp, segments_to_plain_text, speaker_label

DISPLAY_TZ = ZoneInfo("Asia/Ho_Chi_Minh")
SOURCE_LABELS = {"live": "Ghi âm trực tiếp", "upload": "File tải lên"}


def _part_titles(session: SessionRead) -> dict[str, str]:
    return {m.id: m.title for m in session.merge_sources}


def _header_lines(session: SessionRead) -> list[str]:
    created = session.created_at.astimezone(DISPLAY_TZ)
    lines = [
        f"Nguồn: {SOURCE_LABELS.get(session.source, session.source)}",
        f"Thời gian tạo: {created:%d/%m/%Y %H:%M}",
    ]
    if session.duration_ms:
        lines.append(f"Thời lượng: {format_timestamp(session.duration_ms)}")
    if session.original_filename:
        lines.append(f"File gốc: {session.original_filename}")
    if session.group:
        lines.append(f"Nhóm: {session.group.name}")
    if session.merge_sources:
        lines.append(f"Gộp từ {len(session.merge_sources)} phiên: " + "; ".join(m.title for m in session.merge_sources))
    return lines


def build_txt(session: SessionRead) -> bytes:
    parts = [session.title, "=" * max(len(session.title), 10), *_header_lines(session), ""]
    if session.summary:
        s = session.summary
        parts += ["TÓM TẮT", s.summary, ""]
        if s.key_points:
            parts += ["Ý chính:", *[f"- {p}" for p in s.key_points], ""]
        if s.action_items:
            parts.append("Việc cần làm:")
            for item in s.action_items:
                extra = ", ".join(x for x in (item.owner, item.due) if x)
                parts.append(f"- {item.task}" + (f" ({extra})" if extra else ""))
            parts.append("")
        if s.decisions:
            parts += ["Quyết định:", *[f"- {d}" for d in s.decisions], ""]
        parts += ["TRANSCRIPT", ""]
    parts.append(segments_to_plain_text(session.segments, _part_titles(session)))
    # BOM giúp Notepad cũ trên Windows hiển thị đúng tiếng Việt.
    return ("﻿" + "\n".join(parts) + "\n").encode("utf-8")


def build_docx(session: SessionRead) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    doc.add_heading(session.title, level=0)
    meta = doc.add_paragraph()
    for i, line in enumerate(_header_lines(session)):
        run = meta.add_run(("\n" if i else "") + line)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x6B, 0x66, 0x5E)

    if session.summary:
        s = session.summary
        doc.add_heading("Tóm tắt", level=1)
        doc.add_paragraph(s.summary)
        if s.key_points:
            doc.add_heading("Ý chính", level=2)
            for p in s.key_points:
                doc.add_paragraph(p, style="List Bullet")
        if s.action_items:
            doc.add_heading("Việc cần làm", level=2)
            for item in s.action_items:
                para = doc.add_paragraph(style="List Bullet")
                para.add_run(item.task)
                extra = " · ".join(x for x in (item.owner, item.due) if x)
                if extra:
                    r = para.add_run(f" — {extra}")
                    r.italic = True
        if s.decisions:
            doc.add_heading("Quyết định", level=2)
            for d in s.decisions:
                doc.add_paragraph(d, style="List Bullet")

    doc.add_heading("Transcript", level=1)
    part_titles = _part_titles(session)
    current_origin = None
    for seg in session.segments:
        if part_titles and seg.origin and seg.origin != current_origin:
            current_origin = seg.origin
            doc.add_heading(part_titles.get(seg.origin, "Phần gộp"), level=2)
        para = doc.add_paragraph()
        label_parts = [p for p in (speaker_label(seg.speaker), format_timestamp(seg.start_ms)) if p]
        if label_parts:
            label = para.add_run(" · ".join(label_parts) + "\n")
            label.bold = True
            label.font.size = Pt(9)
            label.font.color.rgb = RGBColor(0x2F, 0x5D, 0x50)
        para.add_run(seg.text)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
