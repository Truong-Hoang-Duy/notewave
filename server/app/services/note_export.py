"""Xuất ghi chú Cornell sang file .txt và .docx (GĐ5)."""

import io
from urllib.parse import quote
from zoneinfo import ZoneInfo

from docx import Document
from docx.shared import Pt, RGBColor

from app.models.note import NoteRead
from app.services.markdown_docx import add_markdown

DISPLAY_TZ = ZoneInfo("Asia/Ho_Chi_Minh")


def _header_lines(note: NoteRead) -> list[str]:
    created = note.created_at.astimezone(DISPLAY_TZ)
    updated = note.updated_at.astimezone(DISPLAY_TZ)
    lines = [
        f"Thời gian tạo: {created:%d/%m/%Y %H:%M}",
        f"Lần sửa gần nhất: {updated:%d/%m/%Y %H:%M}",
    ]
    if note.folder:
        lines.append(f"Thư mục: {note.folder.name}")
    if note.tags:
        lines.append("Thẻ: " + ", ".join(f"#{t.name}" for t in note.tags))
    return lines


def build_note_txt(note: NoteRead) -> bytes:
    parts = [
        note.title,
        "=" * max(len(note.title), 10),
        *_header_lines(note),
        "",
    ]
    active_cues = [c for c in note.cues if c.text.strip()]
    if active_cues:
        parts += [
            "--------------------------------------------------",
            "CÂU HỎI & TỪ KHOÁ (Cột gợi ý Cornell)",
            "--------------------------------------------------",
            *[f"{i + 1}. {c.text}" for i, c in enumerate(active_cues)],
            "",
        ]

    parts += [
        "--------------------------------------------------",
        "NỘI DUNG GHI CHÚ",
        "--------------------------------------------------",
        note.content_md or "(Chưa có nội dung)",
        "",
    ]

    if note.summary.strip():
        parts += [
            "--------------------------------------------------",
            "TÓM TẮT CỦA BẠN",
            "--------------------------------------------------",
            note.summary,
            "",
        ]

    if note.ai_summary:
        s = note.ai_summary
        parts += [
            "--------------------------------------------------",
            "TÓM TẮT AI",
            "--------------------------------------------------",
            s.summary,
            "",
        ]
        if s.key_points:
            parts += ["Ý chính:", *[f"- {p}" for p in s.key_points], ""]
        if s.concepts:
            parts += ["Khái niệm cốt lõi:", *[f"- {c.term}: {c.meaning}" for c in s.concepts], ""]
        if s.review_questions:
            parts += ["Câu hỏi ôn tập AI:", *[f"- {q}" for q in s.review_questions], ""]

    # BOM giúp Notepad cũ trên Windows hiển thị đúng tiếng Việt
    return ("\ufeff" + "\n".join(parts) + "\n").encode("utf-8")


def build_note_docx(note: NoteRead) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    doc.add_heading(note.title, level=0)
    meta = doc.add_paragraph()
    for i, line in enumerate(_header_lines(note)):
        run = meta.add_run(("\n" if i else "") + line)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x6B, 0x66, 0x5E)

    active_cues = [c for c in note.cues if c.text.strip()]
    if active_cues:
        doc.add_heading("Câu hỏi & Từ khoá (Cornell)", level=1)
        for c in active_cues:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(c.text)

    doc.add_heading("Nội dung ghi chú", level=1)
    if note.content_md.strip():
        add_markdown(doc, note.content_md)
    else:
        doc.add_paragraph("(Chưa có nội dung)")

    if note.summary.strip():
        doc.add_heading("Tóm tắt của bạn", level=1)
        doc.add_paragraph(note.summary)

    if note.ai_summary:
        s = note.ai_summary
        doc.add_heading("Tóm tắt AI", level=1)
        doc.add_paragraph(s.summary)
        if s.key_points:
            doc.add_heading("Ý chính", level=2)
            for p in s.key_points:
                doc.add_paragraph(p, style="List Bullet")
        if s.concepts:
            doc.add_heading("Khái niệm cốt lõi", level=2)
            for c in s.concepts:
                p = doc.add_paragraph(style="List Bullet")
                r = p.add_run(f"{c.term}: ")
                r.bold = True
                p.add_run(c.meaning)
        if s.review_questions:
            doc.add_heading("Câu hỏi ôn tập", level=2)
            for q in s.review_questions:
                doc.add_paragraph(q, style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
