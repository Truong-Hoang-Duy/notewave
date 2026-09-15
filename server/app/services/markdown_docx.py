"""Chuyển Markdown (kết quả Mistral OCR) sang đoạn văn .docx — đủ cho các cấu trúc OCR hay trả về:
heading, đoạn văn, danh sách (có thụt cấp), trích dẫn, bảng, khối code, in đậm / nghiêng / code inline."""

import re

from docx.document import Document as DocxDocument
from docx.shared import Pt
from docx.text.paragraph import Paragraph

_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
_QUOTE = re.compile(r"^\s*>\s?(.*)$")
_RULE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
_TABLE_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|_[^_\s][^_]*_|`[^`]+`)")
_LINK = re.compile(r"!?\[([^\]]*)\]\([^)]*\)")
_BR = re.compile(r"<br\s*/?>", re.IGNORECASE)


def _add_inline(paragraph: Paragraph, text: str) -> None:
    text = _BR.sub("\n", _LINK.sub(r"\1", text))
    for part in _INLINE.split(text):
        if not part:
            continue
        if (part.startswith("**") and part.endswith("**")) or (part.startswith("__") and part.endswith("__")):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        elif len(part) > 2 and part[0] in "*_" and part[-1] == part[0]:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def _add_table(doc: DocxDocument, lines: list[str]) -> None:
    rows = [_split_row(line) for line in lines if not _TABLE_SEPARATOR.match(line)]
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c in range(width):
            cell = table.cell(r, c)
            paragraph = cell.paragraphs[0]
            _add_inline(paragraph, row[c] if c < len(row) else "")
            if r == 0 and len(lines) > 1 and _TABLE_SEPARATOR.match(lines[1]):
                for run in paragraph.runs:
                    run.bold = True
    doc.add_paragraph()


def add_markdown(doc: DocxDocument, markdown: str) -> None:
    lines = markdown.splitlines()
    i = 0
    paragraph_lines: list[str] = []

    def flush() -> None:
        if paragraph_lines:
            _add_inline(doc.add_paragraph(), "\n".join(paragraph_lines))
            paragraph_lines.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush()
            i += 1
            code: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            run = doc.add_paragraph().add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            i += 1
            continue

        if stripped.startswith("|"):
            flush()
            block: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            _add_table(doc, block)
            continue

        if not stripped or _RULE.match(line):
            flush()
        elif heading := _HEADING.match(line):
            flush()
            # Cấp 0-1 đã dùng cho tiêu đề phiên và "Nội dung tài liệu".
            doc.add_heading(heading.group(2), level=min(len(heading.group(1)) + 1, 9))
        elif bullet := _BULLET.match(line):
            flush()
            style = "List Bullet 2" if len(bullet.group(1).expandtabs(4)) >= 2 else "List Bullet"
            _add_inline(doc.add_paragraph(style=style), bullet.group(2))
        elif numbered := _NUMBERED.match(line):
            flush()
            style = "List Number 2" if len(numbered.group(1).expandtabs(4)) >= 2 else "List Number"
            _add_inline(doc.add_paragraph(style=style), numbered.group(2))
        elif quote := _QUOTE.match(line):
            flush()
            _add_inline(doc.add_paragraph(style="Quote"), quote.group(1))
        else:
            paragraph_lines.append(stripped)
        i += 1
    flush()
